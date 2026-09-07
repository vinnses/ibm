#!/usr/bin/env python3
"""Reproducibly inspect and extract the preserved UFPR CPA workbook.

The source workbook contains dangling drawing/VML relationships which make
openpyxl reject it.  A temporary ZIP copy removes only relationships pointing
to absent package members; source bytes are never changed and the operation is
recorded in the inventory.
"""
from __future__ import annotations

import argparse, csv, gzip, hashlib, json, os, posixpath, tempfile
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile
from xml.etree import ElementTree as ET

from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries

ROOT = Path(__file__).resolve().parents[1]
SOURCE_REL = Path("administracao/historico/fontes/documentos/cpa-avaliacao-curso-informatica-biomedica-2022.xlsx")
SOURCE = ROOT / SOURCE_REL
OUT = ROOT / "administracao/dados/cpa"
EXPECTED_SHA = "4c968ff6e4c32b9d023cfd4f783c0a38b3474bbcde231c383587a9198771cf40"
COURSE_CODE = "40001016064G0"
COURSE_NAME = "INFORMÁTICA BIOMÉDICA"
PERIOD_LITERAL = "AVALIAÇÃO DE CURSOS 2022 (13 a 28 de fevereiro 2023)"
SOURCE_PAGE_REL = Path("administracao/historico/fontes/paginas/cpa-avaliacao-cursos-2022-2026-09-04.html")
SOURCE_PAGE = ROOT / SOURCE_PAGE_REL
SHEETS = ['dados', 'HAB_LETRAS', 'Quantidade de respostas', 'Frequencia de respostas', 'Questoes', 'criterio_analise', 'Setor', 'Cursos', 'Modelo', 'SEGURANÇA', 'AMBIENTES', 'SERVIÇOS TERCEIRIZADOS', 'SALAS DE AULA', 'NTE', 'LABORATÓRIOS', 'BIBLIOTECAS', 'SISTEMA DE REDES', 'AVALIAÇÃO DE SERVIÇOS ', 'INST_PESQUISA ']

def sha256(p: Path) -> str:
    h = hashlib.sha256()
    with p.open('rb') as f:
        for b in iter(lambda: f.read(1024 * 1024), b''): h.update(b)
    return h.hexdigest()

def safe(v):
    if v is None or isinstance(v, (str, int, float, bool)): return v
    return str(v)

def sanitized_copy(source: Path):
    fd, name = tempfile.mkstemp(suffix='.xlsx'); os.close(fd); target = Path(name)
    missing = []
    with ZipFile(source) as zin, ZipFile(target, 'w', ZIP_DEFLATED) as zout:
        names = set(zin.namelist())
        for n in sorted(names):
            b = zin.read(n)
            if n.startswith('xl/worksheets/_rels/') and n.endswith('.rels'):
                root = ET.fromstring(b)
                changed = False
                for rel in list(root):
                    t = rel.attrib.get('Target')
                    target_name = posixpath.normpath(posixpath.dirname(n) + '/../' + t) if t else None
                    rel_type = rel.attrib.get('Type', '')
                    # Do not touch external hyperlinks or any other relationship.
                    # The preserved source has only dangling internal drawing/VML
                    # relationships, which are safe to omit for cell inspection.
                    removable = rel_type.endswith('/drawing') or rel_type.endswith('/vmlDrawing')
                    if target_name and target_name not in names and removable and not rel.attrib.get('TargetMode'):
                        missing.append({'relationship_file': n, 'target': t, 'target_normalized': target_name, 'type': rel.attrib.get('Type')})
                        root.remove(rel); changed = True
                if changed: b = ET.tostring(root, encoding='utf-8', xml_declaration=True)
            zout.writestr(n, b)
    return target, missing

def table_context(ws, row, col):
    contexts = []
    for name in ws.tables:
        table = ws.tables[name]
        min_col, min_row, max_col, max_row = range_boundaries(table.ref)
        if min_row <= row <= max_row and min_col <= col <= max_col:
            header = ws.cell(min_row, col)
            contexts.append({"table": name, "range": table.ref,
                             "header_cell": header.coordinate,
                             "header_value": safe(header.value)})
    return contexts

def merged_context(ws, coord):
    for rng in ws.merged_cells.ranges:
        if coord in rng: return str(rng)
    return None

def cell_record(ws, c, cache_ws):
    formula = c.value if c.data_type == "f" or (isinstance(c.value, str) and c.value.startswith('=')) else None
    cached = cache_ws[c.coordinate].value if formula else None
    return {'worksheet': ws.title, 'cell': c.coordinate, 'row': c.row, 'column': c.column,
            'value': safe(None if formula else c.value), 'formula': formula, 'cached_value': safe(cached),
            'formula_cache_status': ('stored' if cached is not None else 'missing') if formula else None,
            'data_type': c.data_type, 'number_format': c.number_format, 'style_id': c.style_id,
            'table_context': table_context(ws, c.row, c.column), 'merged_range': merged_context(ws, c.coordinate),
            'hyperlink': c.hyperlink.target if c.hyperlink else None,
            'comment': c.comment.text if c.comment else None}

def main() -> int:
    ap = argparse.ArgumentParser(); ap.add_argument('--source', type=Path, default=SOURCE); ap.add_argument('--output', type=Path, default=OUT); args = ap.parse_args()
    source = args.source; output = args.output
    if not source.exists(): raise SystemExit(f'missing source: {source}')
    actual = sha256(source)
    if actual != EXPECTED_SHA: raise SystemExit(f'SHA mismatch: expected {EXPECTED_SHA}, got {actual}')
    output.mkdir(parents=True, exist_ok=True)
    source_label = source.relative_to(ROOT).as_posix() if source.is_relative_to(ROOT) else str(source)
    try:
        direct_workbook = load_workbook(source, data_only=False)
        direct_workbook.close()
        direct_load_error = None
    except Exception as exc:
        direct_load_error = f"{type(exc).__name__}: {exc}"
    temp, missing = sanitized_copy(source)
    try:
        wb_formula = load_workbook(temp, data_only=False)
        wb_cache = load_workbook(temp, data_only=True)
        if wb_formula.sheetnames != SHEETS: raise RuntimeError('worksheet set/order mismatch')
        all_cells=[]; sheet_meta=[]; populated={}
        for ws in wb_formula.worksheets:
            cache_ws=wb_cache[ws.title]; cells=[]; formulas=0
            for row in ws.iter_rows():
                for c in row:
                    if c.value is not None:
                        rec=cell_record(ws,c,cache_ws); cells.append(rec); all_cells.append(rec)
                        formulas += int(bool(rec['formula']))
            populated[ws.title]=cells
            sheet_meta.append({'worksheet':ws.title,'dimensions':ws.calculate_dimension(),'max_row':ws.max_row,'max_column':ws.max_column,
                'populated_cells':len(cells),'formula_cells':formulas,'merged_ranges':sorted(map(str,ws.merged_cells.ranges)),
                'hidden_rows':sorted(i for i,d in ws.row_dimensions.items() if d.hidden),
                'hidden_columns':sorted(i for i,d in ws.column_dimensions.items() if d.hidden),
                'hyperlinks':sum(bool(c.hyperlink) for row in ws.iter_rows() for c in row if c.value is not None),
                'comments':sum(bool(c.comment) for row in ws.iter_rows() for c in row if c.value is not None),
                'tables':[{'name': name, 'range': ws.tables[name].ref,
                           'header_row': range_boundaries(table.ref)[1],
                           'headers': [safe(ws.cell(range_boundaries(table.ref)[1], col).value)
                                       for col in range(range_boundaries(table.ref)[0], range_boundaries(table.ref)[2] + 1)]}
                          for name in sorted(ws.tables) for table in [ws.tables[name]]],
                'header_structure': ('single table header row for each recognized table'
                                     if ws.tables else
                                     'non-tabular/presentation layout; merged ranges and cell coordinates preserve header relationships'),
                'data_types': {k:sum(1 for c in cells if c['data_type']==k) for k in sorted(set(c['data_type'] for c in cells))},
                'number_formats': sorted(set(c['number_format'] for c in cells))})
        inventory={'source':{'path':source_label,'sha256':actual,'size_bytes':source.stat().st_size,'expected_sha256':EXPECTED_SHA,
                             'recorded_hash_match': actual == EXPECTED_SHA,
                             'context_page_path': SOURCE_PAGE_REL.as_posix(),
                             'context_page_sha256': sha256(SOURCE_PAGE)},
            'workbook_properties':{'creator':safe(wb_formula.properties.creator),'title':safe(wb_formula.properties.title),'subject':safe(wb_formula.properties.subject),'created':safe(wb_formula.properties.created),'modified':safe(wb_formula.properties.modified),'calc_mode':safe(wb_formula.calculation.calcMode),'full_calc_on_load':safe(wb_formula.calculation.fullCalcOnLoad)},
            'worksheets':sheet_meta,'worksheet_count':len(sheet_meta),'populated_cell_count':len(all_cells),'formula_cell_count':sum(1 for c in all_cells if c['formula']),
            'formula_cached_value_count':sum(1 for c in all_cells if c['formula'] and c['formula_cache_status']=='stored'),
            'formula_missing_cache_count':sum(1 for c in all_cells if c['formula'] and c['formula_cache_status']=='missing'),
            'merged_range_count':sum(len(sheet['merged_ranges']) for sheet in sheet_meta),
            'hidden_row_count':sum(len(sheet['hidden_rows']) for sheet in sheet_meta),
            'hidden_column_count':sum(len(sheet['hidden_columns']) for sheet in sheet_meta),
            'hyperlink_cell_count':sum(sheet['hyperlinks'] for sheet in sheet_meta),
            'comment_cell_count':sum(sheet['comments'] for sheet in sheet_meta),
            'openpyxl_direct_load_error':direct_load_error,
            'dangling_relationships_removed_for_openpyxl':missing,
            'observations':['Direct openpyxl loading fails because workbook XML points to absent drawing/VML package members. A temporary sanitized ZIP removes only those 25 dangling internal drawing/VML relationships. Source bytes are unchanged.',
                            'No formulas were recalculated. Formula text and the workbook stored cache are separate; a stored cache is preserved but not presumed current.',
                            'The preserved CPA page labels the source as academic-calendar 2022 results. Questoes!C2:C48 states the literal evaluation period 13–28 February 2023; presentation sheets say 2022_2.',
                            'The source quantity table explicitly labels three course respondents and 132 course responses; no denominator beyond source-labelled fields is imputed.']}
        (output/'workbook-inventory.json').write_text(json.dumps(inventory,ensure_ascii=False,indent=2,sort_keys=True)+'\n')
        # The faithful intermediate is compressed to remain below ordinary
        # repository file limits.  mtime=0 and an empty filename make it stable.
        (output/'observed-cells.jsonl').unlink(missing_ok=True)
        with (output/'observed-cells.jsonl.gz').open('wb') as raw, gzip.GzipFile(filename='', mode='wb', fileobj=raw, mtime=0) as f:
            for c in all_cells: f.write((json.dumps(c,ensure_ascii=False,sort_keys=True,allow_nan=False)+'\n').encode('utf-8'))
        prov={'workbook_path':source_label,'workbook_sha256':actual,'worksheet_count':len(SHEETS),
              'context_page_path':SOURCE_PAGE_REL.as_posix(),'context_page_sha256':sha256(SOURCE_PAGE),
              'source_url':'https://cpa.ufpr.br/wp-content/uploads/2023/04/informatica-biomedica.xlsx',
              'accessed_at':'2026-09-04','document_period':'academic calendar 2022; presentation label 2022_2',
              'fieldwork_period_literal':PERIOD_LITERAL,'fieldwork_period_source':'Questoes!C2:C48',
              'course_code':COURSE_CODE,'course_name':COURSE_NAME,
              'respondent_universe_literal':'QTD_RESPONDENTES_CURSO','respondent_count':3,
              'response_count':132,'extraction_status':'source-derived; no imputation'}
        (output/'provenance.json').write_text(json.dumps(prov,ensure_ascii=False,indent=2,sort_keys=True)+'\n')
        def write_csv(name, fields, rows):
            with (output/name).open('w',newline='',encoding='utf-8') as f:
                w=csv.DictWriter(f,fieldnames=fields); w.writeheader(); w.writerows(rows)
        data=populated['dados']; headers=[c['value'] for c in populated['dados'] if c['row']==1]
        # use coordinates/values rather than pandas: source row is retained exactly.
        rowmap={}
        for c in data:
            if c['row']>1: rowmap.setdefault(c['row'],{})[headers[c['column']-1] if c['column']<=len(headers) else f'column_{c["column"]}']=(c['cached_value'] if c['formula'] else c['value'])
        obs=[]
        for r, vals in sorted(rowmap.items()):
            if vals.get('COD_CURSO')==COURSE_CODE:
                obs.append({'source_worksheet':'dados','source_cell':f'A{r}:N{r}','source_row':r,'source_workbook':source_label,'workbook_sha256':actual,'period':PERIOD_LITERAL,'period_source':'Questoes!C2:C48','id_pesquisa':vals.get('ID_PESQUISA'),'questionnaire_id':vals.get('ID_QUESTIONARIO'),'question_id':vals.get('ID_PERGUNTA'),'questionnaire':vals.get('QUESTIONARIO'),'question':vals.get('PERGUNTA'),'course_code':vals.get('COD_CURSO'),'course_literal':vals.get('CURSO'),'sector_code':vals.get('COD_SETOR_CURSO'),'sector_literal':vals.get('SETOR_CURSO'),'response_type':vals.get('TIPO_QUEST2'),'institution':vals.get('Instituição'),'category':vals.get('QUESTIONARIO'),'value':vals.get('RESPOSTA'),'unit':'literal response','universe':'three course respondents as explicitly labelled in Quantidade de respostas','denominator':3,'notes':'One source response row. ID_PESQUISA is retained literally and is not reinterpreted as a person identifier. Course code/name are copied from the same row; respondent count is copied from the course quantity table.','structural_status':'direct source row; no category reinterpretation'})
        fields=list(obs[0]) if obs else ['source_worksheet','source_cell','source_row','source_workbook','workbook_sha256','period','id_pesquisa','questionnaire_id','question_id','questionnaire','question','course_code','course_literal','sector_code','sector_literal','response_type','institution','category','value','unit','universe','denominator','notes','structural_status']
        write_csv('cpa_observations.csv',fields,obs)
        # Source responses are categorical. Preserve literal categories and
        # derive only auditable frequencies within each question; denominator
        # three is explicit in the workbook's course quantity rows.
        freq=[]
        by_question={}
        for item in obs:
            key=(item['questionnaire_id'],item['question_id'],item['questionnaire'],item['question'],item['response_type'])
            by_question.setdefault(key,[]).append(item)
        for key, items in sorted(by_question.items(), key=lambda pair: (int(pair[0][0]), int(pair[0][1]))):
            values={}
            for item in items: values.setdefault(str(item['value']),[]).append(item)
            for value, members in sorted(values.items()):
                freq.append({'source_worksheet':'dados','source_cells':';'.join(f"F{x['source_row']}" for x in members),
                    'source_workbook':source_label,'workbook_sha256':actual,'period':PERIOD_LITERAL,
                    'period_source':'Questoes!C2:C48',
                    'questionnaire_id':key[0],'question_id':key[1],'questionnaire':key[2],'question':key[3],
                    'response_type':key[4],'category':value,'value':len(members),'unit':'responses',
                    'universe':'responses to this question from the three explicitly labelled course respondents',
                    'denominator':len(items),'percentage':len(members)/len(items),
                    'notes':'Count and proportion derived only from cited literal response cells; categories are not combined.',
                    'structural_status':'mechanical frequency from source rows'})
        write_csv('cpa_response_frequencies.csv',list(freq[0]),freq)
        q=populated['Questoes']; qh=[c['value'] for c in q if c['row']==1]; qrows={}
        for c in q:
            if c['row']>1: qrows.setdefault(c['row'],{})[qh[c['column']-1] if c['column']<=len(qh) else f'column_{c["column"]}']=(c['cached_value'] if c['formula'] else c['value'])
        qout=[]
        for r,v in sorted(qrows.items()): qout.append({'source_worksheet':'Questoes','source_cell':f'A{r}:U{r}','source_row':r,'source_workbook':source_label,'workbook_sha256':actual,**{k:v.get(k) for k in qh}})
        write_csv('cpa_questions.csv', ['source_worksheet','source_cell','source_row','source_workbook','workbook_sha256']+qh, qout)
        criteria=populated['criterio_analise']; kh=[c['value'] for c in criteria if c['row']==1]; krows=[]
        for r in range(2, 7):
            vals={kh[c['column']-1] if c['column']<=len(kh) else f'column_{c["column"]}':(c['cached_value'] if c['formula'] else c['value']) for c in criteria if c['row']==r}
            if vals.get('tipo_quest'):
                krows.append({'source_worksheet':'criterio_analise','source_cell':f'A{r}:G{r}','source_row':r,
                              'source_workbook':source_label,'workbook_sha256':actual,**vals,
                              'structural_status':'literal source criterion; not applied or interpreted'})
        write_csv('cpa_response_criteria.csv',['source_worksheet','source_cell','source_row','source_workbook','workbook_sha256']+kh+['structural_status'],krows)
        cr=populated['Quantidade de respostas']; ch=[c['value'] for c in cr if c['row']==1]; crows={}
        for c in cr:
            if c['row']>1: crows.setdefault(c['row'],{})[ch[c['column']-1] if c['column']<=len(ch) else f'column_{c["column"]}'] = (c['cached_value'] if c['formula'] else c['value'])
        cout=[]
        for r,v in sorted(crows.items()):
            if v.get('COD_CURSO')==COURSE_CODE: cout.append({'source_worksheet':'Quantidade de respostas','source_cell':f'A{r}:L{r}','source_row':r,'source_workbook':source_label,'workbook_sha256':actual,**{k:v.get(k) for k in ch},'universe':'respondents as labeled by source','denominator':v.get('QTD_RESPONDENTES_CURSO'),'structural_status':'source-labeled count; formulas represented with their stored workbook cache, without recalculation'})
        write_csv('cpa_course_response_counts.csv',['source_worksheet','source_cell','source_row','source_workbook','workbook_sha256']+ch+['universe','denominator','structural_status'],cout)
        courses=populated['Cursos']; course_headers=[c['value'] for c in courses if c['row']==1]; registry=[]
        for r in range(2, max(c['row'] for c in courses)+1):
            vals={course_headers[c['column']-1] if c['column']<=len(course_headers) else f'column_{c["column"]}':(c['cached_value'] if c['formula'] else c['value']) for c in courses if c['row']==r}
            if vals.get('COD_CURSO')==COURSE_CODE:
                registry.append({'source_worksheet':'Cursos','source_cell':f'A{r}:G{r}','source_row':r,'source_workbook':source_label,'workbook_sha256':actual,**vals,'structural_status':'literal course registry row'})
        write_csv('cpa_course_registry.csv', ['source_worksheet','source_cell','source_row','source_workbook','workbook_sha256']+course_headers+['structural_status'], registry)
        (output/'README.md').write_text(f'''# CPA extraction\n\nSource: `{source_label}`; SHA-256 `{actual}`; title `INFORMÁTICA BIOMÉDICA`; institution Comissão Própria de Avaliação — UFPR; source URL `https://cpa.ufpr.br/wp-content/uploads/2023/04/informatica-biomedica.xlsx`; accessed 2026-09-04; document period academic calendar 2022 (workbook presentation label `2022_2`); document type CPA response workbook. The preserved page is `{SOURCE_PAGE_REL.as_posix()}` with SHA-256 `{sha256(SOURCE_PAGE)}`. The question registry states `{PERIOD_LITERAL}` in `Questoes!C2:C48`; this corrects the older manifest note that the field period was not stated.\n\nInstall the pinned dependency with `python -m pip install -r scripts/requirements-w030-cpa.txt`, then run `python scripts/extract_cpa_workbook.py`. Extraction uses Python and openpyxl 3.1.5. Direct loading fails because the original ZIP points to 25 absent internal drawing/VML members. A temporary ZIP removes only those dangling internal drawing/VML relationships so openpyxl can read cells. External relationships remain, and the original is unchanged. Formulas and stored cached values are separate in `observed-cells.jsonl.gz`; formulas are never recalculated, caches are not presumed current, and missing caches are explicit. Decompress with `gzip -dc observed-cells.jsonl.gz`.\n\n`observed-cells.jsonl.gz` represents every populated cell with coordinate, value/formula/cache, type/format, table header context, merge, hyperlink and comment. CSV outputs are source-derived: `cpa_observations.csv` has 132 literal response rows for three source-labelled respondents and 44 questions; `ID_PESQUISA` is retained literally without assuming that it is a person identifier. `cpa_response_frequencies.csv` mechanically counts the uncombined literal response categories per question and divides only by the three explicit response rows; `cpa_course_registry.csv` proves that code `{COURSE_CODE}` is labelled `{COURSE_NAME}`; `cpa_course_response_counts.csv` preserves the three quantity rows; `cpa_questions.csv` preserves the 47 populated question-registry rows; and `cpa_response_criteria.csv` preserves the five source criterion/type rows without applying them. Three registered questions have no emitted course response row; absence is not converted to zero. Presentation-sheet cached values are not normalized because they contain saved computed displays and some missing/stale caches; formula and cache remain fully auditable in the intermediate layer.\n''')
    finally: temp.unlink(missing_ok=True)
    return 0
if __name__=='__main__':
    try: raise SystemExit(main())
    except Exception as exc:
        raise SystemExit(f'CPA extraction failed: {exc}')
